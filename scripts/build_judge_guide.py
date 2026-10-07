"""Build a concise architecture and judge preparation document from verified scope."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
QA=ROOT/'.runtime/judge-guide';QA.mkdir(exist_ok=True)

def diagram():
    im=Image.new('RGB',(1440,555),'white');draw=ImageDraw.Draw(im)
    bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',28)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',24)
    boxes=[(20,20,420,180,'Log inputs',['Pasted or uploaded logs','Sanitized public demo feed']),
           (520,20,920,230,'Python FastAPI backend',['Normalize and correlate','Workers and explained priority','Validate and cache AI output']),
           (1020,20,1420,180,'SQLite local storage',['Raw evidence and entity links','Incidents context audit reports']),
           (20,360,420,540,'Tier 1 browser workspace',['Analysis and Log Lab','Review actions and report draft']),
           (520,360,920,540,'NVIDIA hosted NIM API',['Nemotron classification','Retrieval planning and briefs']),
           (1020,360,1420,540,'Neo4j MITRE knowledge',['Read only ATT&CK context','Tactics detection mitigation'])]
    def arrow(a,b,two=False):
        draw.line([a,b],fill='#68796f',width=4)
        import math
        angle=math.atan2(b[1]-a[1],b[0]-a[0])
        def tip(p,theta):
            points=[p,(p[0]-15*math.cos(theta-.45),p[1]-15*math.sin(theta-.45)),(p[0]-15*math.cos(theta+.45),p[1]-15*math.sin(theta+.45))]
            draw.polygon(points,fill='#68796f')
        tip(b,angle)
        if two:tip(a,angle+math.pi)
    arrow((420,100),(520,100));arrow((920,100),(1020,100),True)
    arrow((520,210),(320,360),True);arrow((720,230),(720,360),True);arrow((920,210),(1120,360),True)
    for x1,y1,x2,y2,title,lines in boxes:
        draw.rounded_rectangle((x1,y1,x2,y2),radius=13,fill='#f5f7f6',outline='#cfd8d2',width=2)
        draw.text(((x1+x2)/2,y1+23),title,font=bold,fill='black',anchor='mt')
        for i,line in enumerate(lines):draw.text(((x1+x2)/2,y1+72+i*33),line,font=font,fill='#39463f',anchor='mt')
    path=QA/'architecture.png';im.save(path);return path

doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.65);section.bottom_margin=Inches(.65)
section.left_margin=Inches(.7);section.right_margin=Inches(.7)
section.header_distance=Inches(.25);section.footer_distance=Inches(.25)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(11)
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.line_spacing=1.10
doc.styles['Title'].font.size=Pt(24);doc.styles['Title'].paragraph_format.space_after=Pt(9)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Subtitle'].font.size=Pt(11)
doc.styles['Heading 1'].font.size=Pt(18);doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].paragraph_format.space_before=Pt(9);doc.styles['Heading 2'].paragraph_format.space_after=Pt(4)
doc.styles['Heading 3'].font.size=Pt(11);doc.styles['Heading 3'].paragraph_format.space_before=Pt(8);doc.styles['Heading 3'].paragraph_format.space_after=Pt(3)
for name in ['Heading 1','Heading 2','Heading 3']:doc.styles[name].paragraph_format.keep_with_next=True
doc.core_properties.title='Shiftline Architecture and Demo Pitch'
doc.core_properties.subject='Tier 1 SOC architecture demo and judge questions'
doc.core_properties.author='Shiftline'
header=section.header.paragraphs[0];header.text='SHIFTLINE   /   JUDGE PREPARATION';header.style='Normal'
for r in header.runs:r.font.size=Pt(8)
footer=section.footer.paragraphs[0];footer.text='7 October 2026   |   Enterprise oriented MVP   |   Page '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(8)

def p(text,bold=False):
    para=doc.add_paragraph();run=para.add_run(text);run.bold=bold;return para
def h(text,level=2):return doc.add_heading(text,level)
def bullet(text):
    para=doc.add_paragraph(style='List Bullet');para.add_run(text);para.paragraph_format.space_after=Pt(4);return para
def page(title):doc.add_page_break();h(title,1)
def qa(question,answer):h(question,3);p(answer)
def link(text,url):
    para=doc.add_paragraph();hyper=OxmlElement('w:hyperlink')
    rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyper.set(qn('r:id'),rel);r=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'235747');props.append(color);r.append(props);t=OxmlElement('w:t');t.text=text;r.append(t);hyper.append(r);para._p.append(hyper)
    para.paragraph_format.space_after=Pt(3)

doc.add_heading('Shiftline Architecture and Demo Pitch',0)
p('A practical guide for presenting the Tier 1 SOC prototype and answering judge questions.',True)
p('Shiftline turns incoming security logs into prioritized, explainable incident investigations. This guide covers the current architecture, a short spoken pitch, a live demo and the boundaries you should state confidently.')
p('Position it as an enterprise-oriented MVP. The workflow has evidence, business context and human controls; production identity, resilient infrastructure and operational integrations remain future work.')
h('Simple architecture')
doc.add_picture(str(diagram()),width=Inches(7.1))
p('The backend controls graph queries and action requests. NIM receives selected evidence for analysis. Neo4j supplies the existing MITRE knowledge graph. The event relationship layer and analyst history live in SQLite, not in the remote MITRE graph.')
h('How one batch becomes an investigation')
bullet('Retain submitted originals and hashes, normalize entities and correlate related activity by shared identity and time.')
bullet('Use NIM to classify incident evidence and plan MITRE retrieval. The backend runs bounded, parameterized read queries against Neo4j.')
bullet('Validate evidence references, apply an explained asset-aware priority score and cache the brief. The analyst reviews, records a disposition and prepares a handoff.')
h('Where AI and application logic belong')
p('AI interprets behavior, proposes techniques and recommends attention. Application logic controls source preservation, correlation boundaries, score calculation, reference checks, caching and approvals. This keeps model suggestions reviewable.')
h('Prototype operating limits')
p('Bulk input accepts up to 3,000 events and 8 MB. The dashboard currently inspects up to 1,000 incidents, with 20 cards per page; reports cover up to 200 incidents. Two concurrent AI calls and a 16-job queue keep inference bounded. These are implementation limits, not validated throughput guarantees.')

page('Pitch and live demo')
h('A spoken pitch of about 90 seconds')
p('“A Tier 1 SOC analyst needs to decide which alerts deserve attention and explain that decision to the next shift. Shiftline helps by turning incoming logs into a smaller queue of related incidents, with evidence the analyst can inspect.')
p('Our Python backend preserves the original evidence, extracts entities and groups related activity. NVIDIA NIM classifies behavior and plans retrieval from our existing Neo4j MITRE knowledge graph. The analyst sees a short brief, relevant techniques, an event timeline and clear next steps. Priority includes asset criticality, so business impact matters alongside technical severity.')
p('Business context is central. A scan may be authorized maintenance, but we only suggest that explanation when an explicitly approved exception matches its scope. Every decision remains with the analyst.')
p('The current demo uses 100 sanitized public logs and shows 50 incident groups. We can open the source behind a finding, review the case and generate an evidence-linked report. This is an enterprise-oriented MVP: traceability and human control are implemented; authenticated roles, durable infrastructure and live enforcement are the next production steps.”')
h('A three minute walkthrough')
for text in [
 '0:00–0:30  Open Analysis in Demo workspace. Show logs, incident groups and the ranked cards. Say “Grouping reduces investigation items; it does not prove a triage-time reduction.”',
 '0:30–1:20  Open the prepared PowerShell or credential-access incident. Show the brief, risk reason and one timeline event. Click a citation to reveal the exact retained source and hash.',
 '1:20–1:50  Open MITRE and Asset. Explain candidate mapping versus analyst approval, and show which business fields are known or missing. Do not invent scanner approvals.',
 '1:50–2:20  Show Close Benign, Escalate and notes. In the isolated demo only, use your own analyst name and a defensible reason if you demonstrate a decision. Security response buttons record approval requests, not enforcement.',
 '2:20–3:00  Open Analyst Report. Generate or open the prepared draft, point to evidence and unresolved items, then show Edit → Approve → Submit to local outbox. External delivery is not connected.'
]:p(text)
h('Prepare before the judges arrive')
p('Keep the server running, select Demo workspace and confirm Settings shows the configured NIM API and connected graph. Keep a completed brief and report ready. Use Inject batch now to show incoming data; Auto Demo schedules eight batches every 30 minutes. Keep a downloaded report available if the network is unreliable. Log Lab is an optional live AI example, not the critical path of the timed demo.')

page('Enterprise controls and production boundaries')
p('Say “The MVP includes enterprise-relevant controls.” Be precise about what those controls protect and what still needs engineering before customer deployment.')
table=doc.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
widths=[1.30,2.90,2.90]
for c,w in zip(table.columns,widths):c.width=Inches(w)
for cell,title in zip(table.rows[0].cells,['Control','Implemented today','Production requirement']):cell.text=title
rows=[
 ('Evidence','Original input and hashes; source spans; recorded analysis history.','Protected retention, backups and independently anchored audit evidence.'),
 ('AI grounding','NIM plans bounded retrieval; references and retrieved MITRE IDs checked; unsupported portfolio drafts withheld.','Evaluate semantic accuracy and incident-level recall on labeled data; monitor drift.'),
 ('Human control','Dispositions, mapping approval, revision checks and separate named response approver.','SSO and role-based authorization. Names are currently self-declared.'),
 ('Business context','Asset criticality and explicit source, host, activity, command and time-scoped exceptions.','Verified CMDB, identity, maintenance and change-management integrations.'),
 ('Data separation','Separate local My logs, Demo and Log Lab databases; demo reset archives its database.','Authenticated tenant isolation, encryption and audited access controls.'),
 ('Response and handoff','Contain/block/disable requests only; editable reports and local outbox.','Approved EDR, firewall, identity and ticketing connectors with execution verification.'),
 ('Reliability','Two background workers, bounded AI concurrency, pagination and cached results.','Durable queue, service monitoring, recovery, load tests and agreed service objectives.')]
for row in rows:
    cells=table.add_row().cells
    for cell,text in zip(cells,row):cell.text=text
for ri,row in enumerate(table.rows):
    trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
    if ri==0:
        repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    for ci,cell in enumerate(row.cells):
        cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
        for side in ['top','left','bottom','right']:
            v=OxmlElement('w:'+side);v.set(qn('w:w'),'95');v.set(qn('w:type'),'dxa');margins.append(v)
        tcpr.append(margins);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'404040' if ri==0 else 'f4f6f5' if ri%2 else 'ffffff');tcpr.append(shade)
        borders=OxmlElement('w:tcBorders')
        for side in ['top','left','bottom','right']:
            b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
        tcpr.append(borders)
        for para in cell.paragraphs:
            para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.06
            for run in para.runs:run.font.size=Pt(10);run.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
h('What success would mean')
p('Measure MTTT as first analyst decision time minus incident ingestion time, including queue wait and processing. Measure MTTR separately at explicit closure. Compare equivalent manual and assisted workloads, report sample size and decision quality, and count missed attacks and mistaken benign closures. Faster decisions alone are not sufficient.')
p('The 100-log demo grouped into 50 incidents. That is a snapshot of work-item grouping, not a measured productivity gain or a validated detection-accuracy result.')

page('Cross questions about the architecture')
qa('Why use AI if correlation and scoring have rules','AI handles behavior interpretation, retrieval hypotheses, severity assessment and recommended attention. Deterministic controls preserve evidence and keep grouping and priority explainable. We use a hybrid architecture because a language model should not control database queries or security actions.')
qa('Why Neo4j and SQLite together','Neo4j holds the existing MITRE knowledge graph and related tactic, detection and mitigation context. SQLite stores events, entity links, incident state and analyst history. Neo4j adds relationships for retrieval; it is not the current raw-log store or a full event GraphRAG platform.')
qa('What does NVIDIA NIM actually do','The hosted NIM endpoint runs Nemotron 3.5 Lightning through an OpenAI-compatible client. It produces batch classifications, retrieval plans, briefs and analytics drafts. The backend validates structured output and references before caching accepted results. We are calling hosted inference, not self-hosting a NIM deployment.')
qa('Are all incoming logs fully analyzed by the model','Every accepted log is retained and normalized. Batch classification covers up to 120 incidents per job, in groups of 10, using up to four representative events each. Detailed briefs use up to three distinct original logs per selected incident; initial batch processing prepares four priority briefs. This is not exhaustive model analysis of every event.')
qa('How do you avoid merging unrelated activity','Correlation uses shared host/user, host/source IP, user/source IP or host/process identity within a one-hour span. Scan-source linking across victims has a ten-minute window. Missing identity or time limits grouping; reviewed cases do not silently absorb new evidence. Correlation remains a hypothesis, not proof of a single attacker.')
qa('Can the AI hallucinate even with citations','Yes. Exact references establish traceability, not semantic truth. Schema checks, retrieved-ID checks, an AI evidence-verification pass and human review reduce risk but do not guarantee correctness. If a speculative portfolio draft is withheld, the desk can show cached source-linked NIM investigation recommendations.')
qa('Is the confidence percentage calibrated','No. Confidence is qualitative: low, medium or high. Risk is an explained priority score, not a probability of compromise. Base priority combines severity and the highest registered affected-asset criticality, with visible context adjustments. Alert count does not directly raise risk.')
qa('Does it replace a SIEM or the analyst','The target role is a triage and handoff layer over incoming security evidence. The MVP accepts pasted or uploaded JSON, NDJSON and XML, with a bulk API. It does not replace enterprise collection, detection coverage or a SIEM, and it does not autonomously decide that an attack is confirmed.')

page('Cross questions about readiness and value')
qa('Why is the demo credible','The demo replays public Splunk samples with source provenance and transformed identifiers and times. It includes authentication, endpoint, scan, PowerShell, DNS, web and cloud activity. No analyst dispositions or confirmed-compromise outcomes are fabricated.')
qa('Can you prove faster triage or fewer false positives','Not yet. A pilot must compare manual and assisted MTTT on equivalent labeled incidents, measuring decision quality and mistaken benign closures. Earlier 3,000-event ingestion measurements exclude AI analysis and cannot establish production throughput.')
qa('What happens if NIM or Neo4j is unavailable','Original evidence and cached results remain available. New AI analysis may fail and require retry. The current queue and demo clock are in-process components; durable scheduling and recovery remain production work.')
qa('Can it safely recognize a benign scanner','Only when an explicitly approved exception matches source, targets, activity and time. Follow-on credential-access or log-clearing evidence can block the suggestion. Missing context remains unknown, and benign closure requires analyst review.')
qa('Where do customer logs and secrets go','Credentials stay server-side; selected logs go to hosted NVIDIA inference. Demo sanitization is pattern-based, and pasted customer logs are not automatically redacted. A pilot needs approved data handling, redaction, retention and secret management.')
qa('Can it contain a host or deliver a senior report','Containment, blocking and account-disable buttons record approval requests without enforcement. Reports combine evidence, cached AI observations and actions, with editing, approval and local outbox submission. External execution and delivery need connectors.')
qa('What makes it enterprise ready','The defensible claim is enterprise-oriented MVP with asset-aware priority, business context, traceability and human decisions. Production needs authenticated roles, secure tenancy, protected retention, durable workers, integration validation and measured operating performance. No certification is claimed.')
qa('What is the differentiator and next milestone','The proposed differentiator is explainable triage with business context and source-linked handoff. Next is a controlled SOC pilot measuring MTTT and decision quality, followed by integrations. Competitive advantage and savings still need validation.')
h('Evidence you can cite in the room')
p('Read-only snapshot on 7 October 2026: 100 demo logs, 50 groups, 46 stored NIM classifications, four detailed briefs and zero analyst dispositions. Counts change with replay and review. Scope follows the current app modules; earlier local measurements are recorded in VERIFICATION.md.')
link('Splunk Attack Data public source repository','https://github.com/splunk/attack_data')
link('MITRE ATT&CK knowledge base','https://attack.mitre.org/')
link('NVIDIA Nemotron 3.5 Lightning documentation','https://docs.nvidia.com/nim/large-language-models/2.0.10/get-started/advanced/get-started-nemotron-3.5-lightning.html')

path=OUT/'Shiftline_Architecture_and_Demo_Pitch.docx';doc.save(path);print(path.name)
